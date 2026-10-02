from __future__ import annotations

from io import BytesIO
import unittest

from openpyxl import load_workbook

from app.services.export_cost_model import CostModelExportError, build_cost_model_workbook


def _line(sku: str, **values) -> dict:
    return {"sku": sku, "name": sku.title(), "unit": "un", "auxiliary": False, "price": 1000, "estimate": 2, "historic": None,
            "grade": None, "source": "Estimada", "quantity": 2, "note": None, "partial": False, **values}


class CostModelWorkbookTests(unittest.TestCase):
    def setUp(self) -> None:
        self.payload = {
            "version": 2, "generated_at": "2026-09-27T12:00:00Z", "range": {"start": "2026-06-01", "end": "2026-08-31"},
            "target_houses": 4,
            "subtypes": [
                {"id": 8, "name": "Normal", "houses": 3, "lines": [_line("SMALL"), _line("BIG", quantity=10), _line("NOPRICE", price=None)]},
                {"id": 9, "name": "Espejada", "houses": 1, "lines": [_line("PART", partial=True, note="Parcial: alguna instancia sin cantidad")]},
            ],
            "selection": [{"name": "Normal", "weight": 0.75}, {"name": "Espejada", "weight": 0.25}],
            "extras": [{"sku": "EXTRA", "name": "=Cinta", "unit": "un", "quantity_per_house": 2, "unit_cost": 500, "value_per_house": 1000,
                        "replaces_sku": None, "origin": "Período"}],
            "detail": [
                {"subtype": "Espejada", "category": "10. Techo", "instance": "T1", "sku": "PART", "name": "Part", "unit": "un", "quantity": None, "price": 1000},
                {"subtype": "Normal", "category": "2. Muros", "instance": "M1", "sku": "BIG", "name": "Big", "unit": "un", "quantity": 10, "price": 1000},
            ],
        }

    def build(self):
        output = BytesIO()
        build_cost_model_workbook("Proyecto", self.payload, output)
        output.seek(0)
        return load_workbook(output)

    def test_sheets_follow_the_page(self) -> None:
        workbook = self.build()
        self.assertEqual(workbook.sheetnames, ["Resumen", "Presupuesto Normal", "Presupuesto Espejada", "Fuera de presupuesto", "Detalle ficha"])

    def test_budget_per_house_adds_the_bom_and_the_extras(self) -> None:
        sheet = self.build()["Presupuesto Normal"]
        # Sorted by budgeted cost, materials without price last.
        self.assertEqual([sheet.cell(row=row, column=2).value for row in range(6, 9)], ["BIG", "SMALL", "NOPRICE"])
        labels = {sheet.cell(row=row, column=1).value: row for row in range(9, sheet.max_row + 1)}
        total, extras, final = labels["Materiales de la ficha"], labels["Fuera de presupuesto incluidos (hoja «Fuera de presupuesto»)"], labels["Presupuesto / vivienda"]
        self.assertEqual(sheet.cell(row=total, column=11).value, "=SUM(K6:K8)")
        self.assertEqual(sheet.cell(row=extras, column=11).value, "='Fuera de presupuesto'!F7")
        self.assertEqual(sheet.cell(row=final, column=11).value, f"=K{total}+K{extras}")
        self.assertEqual(sheet["K2"].value, f"=K{final}")

    def test_summary_weights_the_period_mix_and_the_selection(self) -> None:
        summary = self.build()["Resumen"]
        self.assertEqual(summary["A6"].value, "Normal")
        self.assertEqual(summary["D6"].value, "='Presupuesto Normal'!K10")
        self.assertEqual(summary["F6"].value, "=D6+E6")
        self.assertEqual(summary["G6"].value, 1)
        self.assertEqual(summary["G7"].value, 1)
        self.assertEqual(summary["A8"].value, "Proyecto (mezcla del período)")
        self.assertEqual(summary["F8"].value, "=SUMPRODUCT($B6:$B7,F6:F7)/$B8")
        self.assertEqual(summary["F9"].value, "=0.75*F6+0.25*F7")

    def test_title_names_the_quantity_basis(self) -> None:
        self.assertEqual(self.build()["Resumen"]["A1"].value, "Modelo de costos · Proyecto")
        self.payload["basis"] = "Fábrica + obra"
        self.assertEqual(self.build()["Resumen"]["A1"].value, "Modelo de costos · Proyecto · Cantidad Fábrica + obra")

    def test_text_never_becomes_a_formula(self) -> None:
        self.assertEqual(self.build()["Fuera de presupuesto"]["A5"].value, " =Cinta")

    def test_detail_sorts_categories_by_number(self) -> None:
        detail = self.build()["Detalle ficha"]
        self.assertEqual([detail.cell(row=row, column=2).value for row in (5, 6)], ["2. Muros", "10. Techo"])
        self.assertEqual(detail["J6"].value, "Sin cantidad en la ficha")

    def test_a_single_general_subtype_has_no_mix(self) -> None:
        self.payload.update(subtypes=[{"id": None, "name": "General", "houses": 4, "lines": [_line("A")]}], selection=None, extras=[], detail=[])
        workbook = self.build()
        self.assertEqual(workbook.sheetnames, ["Resumen", "Presupuesto"])
        self.assertEqual(workbook["Presupuesto"]["K9"].value, 0)
        self.assertNotIn("Proyecto (mezcla del período)", [cell.value for cell in workbook["Resumen"]["A"]])

    def test_rejects_a_request_without_the_page_snapshot(self) -> None:
        with self.assertRaises(CostModelExportError):
            build_cost_model_workbook("Proyecto", {}, BytesIO())


if __name__ == "__main__":
    unittest.main()
