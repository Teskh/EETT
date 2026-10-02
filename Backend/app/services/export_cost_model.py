"""The cost model workbook: a snapshot of what the cost model page shows, laid
out for reading. The page sends its own lines, prices, mix and materials
outside the BOM, so every total here is the page's total, and building it
needs neither the project's BOM nor the ERP."""
from __future__ import annotations

from datetime import datetime
from math import isfinite
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

MONEY = '"$"#,##0;-"$"#,##0'
QUANTITY = "General"
PERCENT = "0%"

DARK = PatternFill(fill_type="solid", start_color="18181B", end_color="18181B")
LIGHT = PatternFill(fill_type="solid", start_color="F4F4F5", end_color="F4F4F5")
BODY = Font(name="Calibri", size=10)
BOLD = Font(name="Calibri", size=10, bold=True)
HEADER = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
TITLE = Font(name="Calibri", size=14, bold=True)
SUBTITLE = Font(name="Calibri", size=10, color="71717A")
WARNING = Font(name="Calibri", size=10, color="B45309")
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")
WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
TOP_RULE = Border(top=Side(style="thin", color="A1A1AA"))

MAX_SUBTYPES = 200
MAX_ROWS = 100_000


class CostModelExportError(ValueError):
    pass


def validate_cost_model_payload(payload: dict[str, Any] | None) -> None:
    """The workbook is the page's snapshot: without it there is nothing to lay out."""
    subtypes = (payload or {}).get("subtypes")
    if (payload or {}).get("version") != 2 or not isinstance(subtypes, list) or not subtypes:
        raise CostModelExportError("La exportación usa los datos de la página. Recarga el Modelo de costos e intenta de nuevo.")
    if len(subtypes) > MAX_SUBTYPES or sum(len(item.get("lines") or []) for item in subtypes if isinstance(item, dict)) > MAX_ROWS \
            or len((payload or {}).get("detail") or []) > MAX_ROWS:
        raise CostModelExportError("La exportación es demasiado grande.")


def build_cost_model_workbook(project_name: str, payload: dict[str, Any], output: Any) -> None:
    validate_cost_model_payload(payload)
    subtypes = [_subtype(item) for item in payload["subtypes"] if isinstance(item, dict)]
    extras = [_extra(item) for item in payload.get("extras") or [] if isinstance(item, dict)]
    detail = [_detail(item) for item in payload.get("detail") or [] if isinstance(item, dict)]
    generated = _generated_label(payload.get("generated_at"))

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    used = {"Resumen"}

    subtype_sheets = []
    for subtype in subtypes:
        title = _sheet_title("Presupuesto" if len(subtypes) == 1 and subtype["name"] == "General" else f"Presupuesto {subtype['name']}", used)
        subtype_sheets.append(workbook.create_sheet(title))

    extras_total = None
    if extras:
        extras_sheet = workbook.create_sheet(_sheet_title("Fuera de presupuesto", used))
        extras_total = _ref(extras_sheet.title, f"F{_populate_extras(extras_sheet, extras)}")

    subtype_totals = [
        _populate_subtype(sheet, subtype, extras_total=extras_total, generated=generated)
        for sheet, subtype in zip(subtype_sheets, subtypes)
    ]
    if detail:
        _populate_detail(workbook.create_sheet(_sheet_title("Detalle ficha", used)), detail, [item["name"] for item in subtypes])

    _populate_summary(summary, project_name, payload, subtypes, subtype_totals, extras_total=extras_total, generated=generated)
    workbook.properties.title = f"{project_name} - Modelo de costos"
    workbook.save(output)


# --- Sheets -----------------------------------------------------------------

def _populate_summary(ws: Worksheet, project_name: str, payload: dict[str, Any], subtypes: list[dict[str, Any]],
                      totals: list[dict[str, Any]], *, extras_total: str | None, generated: str) -> None:
    ws.sheet_view.showGridLines = False
    basis = _text(payload.get("basis")) if payload.get("basis") else None
    _write(ws, 1, 1, f"Modelo de costos · {project_name}" + (f" · Cantidad {basis}" if basis else ""), font=TITLE)
    _write(ws, 2, 1, f"Generado el {generated} con los datos y precios que mostraba la página.", font=SUBTITLE)
    range_ = payload.get("range") if isinstance(payload.get("range"), dict) else None
    if range_:
        houses = _number(payload.get("target_houses"))
        _write(ws, 3, 1, f"Referencia histórica: {_text(range_.get('start'))} a {_text(range_.get('end'))}"
               + (f" · {int(houses)} viviendas del proyecto" if houses is not None else ""), font=SUBTITLE)

    headers = ["Subtipología", "Viviendas en el período", "Estimado / viv.", "Materiales de la ficha / viv.",
               "Fuera de presupuesto / viv.", "Presupuesto / vivienda", "Materiales incompletos"]
    _header(ws, 5, headers, {2, 3, 4, 5, 6, 7})
    first = 6
    rows_by_name: dict[str, int] = {}
    for offset, (subtype, total) in enumerate(zip(subtypes, totals)):
        row = first + offset
        rows_by_name.setdefault(subtype["name"], row)
        _write(ws, row, 1, subtype["name"], font=BODY)
        _write(ws, row, 2, subtype["houses"], font=BODY, align=RIGHT)
        _write(ws, row, 3, f"={total['estimated']}", font=BODY, align=RIGHT, number_format=MONEY)
        _write(ws, row, 4, f"={total['budget']}", font=BODY, align=RIGHT, number_format=MONEY)
        _write(ws, row, 5, f"={extras_total}" if extras_total else 0, font=BODY, align=RIGHT, number_format=MONEY)
        _write(ws, row, 6, f"=D{row}+E{row}", font=BOLD, align=RIGHT, number_format=MONEY, fill=LIGHT)
        _write(ws, row, 7, total["incomplete"] or None, font=WARNING if total["incomplete"] else BODY, align=RIGHT)
    last = first + len(subtypes) - 1
    row = last + 1

    # The project's house: each subtype weighted by the houses it started in the period.
    if len(subtypes) > 1 and all(item["houses"] is not None for item in subtypes) and sum(item["houses"] for item in subtypes) > 0:
        _write(ws, row, 1, "Proyecto (mezcla del período)", font=HEADER, fill=DARK)
        _write(ws, row, 2, f"=SUM(B{first}:B{last})", font=HEADER, align=RIGHT, fill=DARK)
        for column in "CDEF":
            _write(ws, row, " CDEF".index(column) + 2, f"=SUMPRODUCT($B{first}:$B{last},{column}{first}:{column}{last})/$B{row}",
                   font=HEADER, align=RIGHT, number_format=MONEY, fill=DARK)
        _write(ws, row, 7, None, fill=DARK)
        row += 1

    selection = [item for item in payload.get("selection") or [] if isinstance(item, dict) and _text(item.get("name")) in rows_by_name]
    if len(selection) > 1:
        label = " · ".join(f"{_text(item['name'])} {round((_number(item.get('weight')) or 0) * 100)}%" for item in selection)
        _write(ws, row, 1, f"Selección en pantalla: {label}", font=BOLD, fill=LIGHT)
        for column in "CDEF":
            terms = "+".join(f"{round(_number(item.get('weight')) or 0, 6)}*{column}{rows_by_name[_text(item['name'])]}" for item in selection)
            _write(ws, row, " CDEF".index(column) + 2, f"={terms}", font=BOLD, align=RIGHT, number_format=MONEY, fill=LIGHT)
        row += 1

    notes = [
        "Presupuesto / vivienda: materiales de la ficha con la cantidad elegida en el modelo (estimada, histórica o manual) "
        "más los materiales fuera de presupuesto incluidos. Es el valor que muestra la página.",
        "Proyecto: promedio de las subtipologías, ponderado por las viviendas que inició cada una en el período.",
        "Materiales incompletos: sin precio o con cantidades en blanco en la ficha. Si solo algunas instancias están en blanco, "
        "se cuentan las que tienen cantidad.",
        "Cada hoja «Presupuesto» detalla una subtipología. «Detalle ficha» muestra las cantidades de la ficha por instancia, antes de ajustes.",
    ]
    for offset, note in enumerate(notes, start=row + 1):
        _write(ws, offset, 1, note, font=SUBTITLE)
    _widths(ws, {"A": 44, "B": 14, "C": 16, "D": 18, "E": 18, "F": 18, "G": 14})
    ws.row_dimensions[5].height = 30
    ws.freeze_panes = "B6"


def _populate_subtype(ws: Worksheet, subtype: dict[str, Any], *, extras_total: str | None, generated: str) -> dict[str, Any]:
    ws.sheet_view.showGridLines = False
    headers = ["Material", "SKU", "Unidad", "Precio unit.", "Estimada / viv.", "Costo estimado", "Histórica / viv.",
               "Confianza", "Fuente", "Presup. / viv.", "Costo presup.", "Observación"]
    _write(ws, 1, 1, f"Presupuesto por vivienda · {subtype['name']}", font=TITLE)
    _write(ws, 3, 1, f"Cantidades por vivienda. Precios que mostraba la página el {generated}. Ordenado por costo presupuestado.", font=SUBTITLE)
    _header(ws, 5, headers, {4, 5, 6, 7, 10, 11})

    lines = sorted(subtype["lines"], key=lambda line: (-_cost(line["quantity"], line["price"]) if _cost(line["quantity"], line["price"]) is not None else float("inf"), line["name"].lower()))
    first = 6
    for row, line in enumerate(lines, start=first):
        note = "; ".join(item for item in ("Material auxiliar" if line["auxiliary"] else None, line["note"]) if item)
        values = [line["name"], line["sku"], line["unit"], line["price"], line["estimate"], f'=IF(OR(D{row}="",E{row}=""),"",D{row}*E{row})',
                  line["historic"], line["grade"], line["source"], line["quantity"], f'=IF(OR(D{row}="",J{row}=""),"",D{row}*J{row})', note or None]
        for column, value in enumerate(values, start=1):
            _write(ws, row, column, value, font=WARNING if column == 12 else BOLD if column == 11 else BODY,
                   align=RIGHT if column in {4, 5, 6, 7, 10, 11} else LEFT,
                   number_format=MONEY if column in {4, 6, 11} else QUANTITY if column in {5, 7, 10} else None)
    last = max(first, first + len(lines) - 1)

    total = last + 2
    _total_row(ws, total, "Materiales de la ficha", {6: f"=SUM(F{first}:F{last})", 11: f"=SUM(K{first}:K{last})"}, fill=LIGHT, font=BOLD)
    _total_row(ws, total + 1, "Fuera de presupuesto incluidos (hoja «Fuera de presupuesto»)", {11: f"={extras_total}" if extras_total else 0}, fill=None, font=BODY)
    _total_row(ws, total + 2, "Presupuesto / vivienda", {11: f"=K{total}+K{total + 1}"}, fill=DARK, font=HEADER)
    _write(ws, 2, 1, "Presupuesto / vivienda", font=BOLD)
    _write(ws, 2, 11, f"=K{total + 2}", font=BOLD, align=RIGHT, number_format=MONEY)

    ws.freeze_panes = f"B{first}"
    ws.auto_filter.ref = f"A5:L{last}"
    _widths(ws, {"A": 40, "B": 13, "C": 8, "D": 12, "E": 12, "F": 14, "G": 12, "H": 10, "I": 14, "J": 12, "K": 14, "L": 38})
    ws.row_dimensions[5].height = 30
    incomplete = sum(1 for line in lines if line["price"] is None or line["estimate"] is None or line["partial"])
    return {"estimated": _ref(ws.title, f"F{total}"), "budget": _ref(ws.title, f"K{total}"), "incomplete": incomplete}


def _populate_extras(ws: Worksheet, extras: list[dict[str, Any]]) -> int:
    """Materials withdrawn that no budget explains, included in every subtype's budget per house."""
    ws.sheet_view.showGridLines = False
    _write(ws, 1, 1, "Fuera de presupuesto", font=TITLE)
    _write(ws, 2, 1, "Materiales retirados que ninguna ficha explica, incluidos en el presupuesto de cada vivienda.", font=SUBTITLE)
    _header(ws, 4, ["Material", "SKU", "Unidad", "Q / vivienda", "Costo unitario", "Costo / vivienda", "Reemplaza a", "Origen"], {4, 5, 6})
    first = 5
    extras = sorted(extras, key=lambda item: -(item["value_per_house"] or 0))
    for row, extra in enumerate(extras, start=first):
        cost = f"=D{row}*E{row}" if extra["quantity_per_house"] is not None and extra["unit_cost"] is not None else extra["value_per_house"]
        values = [extra["name"], extra["sku"], extra["unit"], extra["quantity_per_house"], extra["unit_cost"], cost, extra["replaces_sku"], extra["origin"]]
        for column, value in enumerate(values, start=1):
            _write(ws, row, column, value, font=BODY, align=RIGHT if column in {4, 5, 6} else LEFT,
                   number_format=MONEY if column in {5, 6} else QUANTITY if column == 4 else None)
    last = first + len(extras) - 1
    total = last + 2
    _total_row(ws, total, "Total por vivienda", {6: f"=SUM(F{first}:F{last})"}, fill=DARK, font=HEADER, last_column=8)
    ws.freeze_panes = f"B{first}"
    _widths(ws, {"A": 40, "B": 13, "C": 8, "D": 12, "E": 14, "F": 16, "G": 14, "H": 30})
    return total


def _populate_detail(ws: Worksheet, detail: list[dict[str, Any]], subtype_order: list[str]) -> None:
    ws.sheet_view.showGridLines = False
    _write(ws, 1, 1, "Detalle de la ficha por instancia", font=TITLE)
    _write(ws, 2, 1, "Cantidades estimadas de la ficha, antes de ajustes del modelo. Filtra por subtipología para ver una vivienda.", font=SUBTITLE)
    _header(ws, 4, ["Subtipología", "Categoría", "Instancia", "Material", "SKU", "Unidad", "Q / viv.", "Precio unit.", "Costo", "Observación"], {7, 8, 9})
    order = {name: index for index, name in enumerate(subtype_order)}
    rows = sorted(detail, key=lambda item: (order.get(item["subtype"], len(order)), _natural(item["category"]), _natural(item["instance"]), item["name"].lower()))
    first = 5
    for row, item in enumerate(rows, start=first):
        values = [item["subtype"], item["category"], item["instance"], item["name"], item["sku"], item["unit"], item["quantity"], item["price"],
                  f'=IF(OR(G{row}="",H{row}=""),"",G{row}*H{row})', "Sin cantidad en la ficha" if item["quantity"] is None else None]
        for column, value in enumerate(values, start=1):
            _write(ws, row, column, value, font=WARNING if column == 10 else BODY, align=RIGHT if column in {7, 8, 9} else LEFT,
                   number_format=MONEY if column in {8, 9} else QUANTITY if column == 7 else None)
    ws.freeze_panes = f"A{first}"
    ws.auto_filter.ref = f"A4:J{max(first, first + len(rows) - 1)}"
    _widths(ws, {"A": 22, "B": 24, "C": 26, "D": 40, "E": 13, "F": 8, "G": 10, "H": 12, "I": 14, "J": 24})


# --- Payload ----------------------------------------------------------------

def _subtype(item: dict[str, Any]) -> dict[str, Any]:
    houses = _number(item.get("houses"))
    return {
        "name": _text(item.get("name")) or "General",
        "houses": int(houses) if houses is not None else None,
        "lines": [_line(line) for line in item.get("lines") or [] if isinstance(line, dict)],
    }


def _line(item: dict[str, Any]) -> dict[str, Any]:
    price = _number(item.get("price"))
    return {
        "sku": _text(item.get("sku")), "name": _text(item.get("name")) or _text(item.get("sku")), "unit": _text(item.get("unit")),
        "auxiliary": bool(item.get("auxiliary")), "price": price if price is not None and price > 0 else None,
        "estimate": _number(item.get("estimate")), "historic": _number(item.get("historic")), "grade": _text(item.get("grade")) or None,
        "source": _text(item.get("source")), "quantity": _number(item.get("quantity")), "note": _text(item.get("note")) or None,
        "partial": bool(item.get("partial")),
    }


def _extra(item: dict[str, Any]) -> dict[str, Any]:
    return {
        "sku": _text(item.get("sku")), "name": _text(item.get("name")) or _text(item.get("sku")), "unit": _text(item.get("unit")),
        "quantity_per_house": _number(item.get("quantity_per_house")), "unit_cost": _number(item.get("unit_cost")),
        "value_per_house": _number(item.get("value_per_house")), "replaces_sku": _text(item.get("replaces_sku")) or None,
        "origin": _text(item.get("origin")),
    }


def _detail(item: dict[str, Any]) -> dict[str, Any]:
    price = _number(item.get("price"))
    return {
        "subtype": _text(item.get("subtype")) or "General", "category": _text(item.get("category")), "instance": _text(item.get("instance")),
        "sku": _text(item.get("sku")), "name": _text(item.get("name")) or _text(item.get("sku")), "unit": _text(item.get("unit")),
        "quantity": _number(item.get("quantity")), "price": price if price is not None and price > 0 else None,
    }


def _number(value: Any) -> float | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return float(value) if isfinite(value) else None


def _text(value: Any, limit: int = 300) -> str:
    if value is None:
        return ""
    text = str(value).strip()[:limit]
    # openpyxl writes a leading "=" as a formula.
    return f" {text}" if text.startswith("=") else text


def _cost(quantity: float | None, price: float | None) -> float | None:
    return None if quantity is None or price is None else quantity * price


def _generated_label(value: Any) -> str:
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00")).astimezone().strftime("%d-%m-%Y %H:%M")
    except ValueError:
        return datetime.now().strftime("%d-%m-%Y %H:%M")


# --- Layout helpers ---------------------------------------------------------

def _write(ws: Worksheet, row: int, column: int, value: Any, *, font: Font | None = None, align: Alignment | None = None,
           number_format: str | None = None, fill: PatternFill | None = None) -> None:
    cell = ws.cell(row=row, column=column, value=value)
    if font is not None:
        cell.font = font
    if align is not None:
        cell.alignment = align
    if number_format is not None:
        cell.number_format = number_format
    if fill is not None:
        cell.fill = fill


def _header(ws: Worksheet, row: int, headers: list[str], numeric: set[int]) -> None:
    for column, label in enumerate(headers, start=1):
        _write(ws, row, column, label, font=HEADER, fill=DARK,
               align=Alignment(horizontal="right" if column in numeric else "left", vertical="center", wrap_text=True))


def _total_row(ws: Worksheet, row: int, label: str, values: dict[int, Any], *, fill: PatternFill | None, font: Font, last_column: int = 12) -> None:
    for column in range(1, last_column + 1):
        cell = ws.cell(row=row, column=column)
        cell.border = TOP_RULE
        if fill is not None:
            cell.fill = fill
    _write(ws, row, 1, label, font=font)
    for column, value in values.items():
        _write(ws, row, column, value, font=font, align=RIGHT, number_format=MONEY)


def _widths(ws: Worksheet, widths: dict[str, float]) -> None:
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width


def _sheet_title(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", " ", name).strip().strip("'")[:31] or "Hoja"
    title, index = base, 2
    while title.lower() in {item.lower() for item in used}:
        suffix = f" ({index})"
        title, index = f"{base[:31 - len(suffix)]}{suffix}", index + 1
    used.add(title)
    return title


def _ref(sheet: str, cell: str) -> str:
    return "'{}'!{}".format(sheet.replace("'", "''"), cell)


def _natural(value: str) -> list[Any]:
    return [(0, int(part)) if part.isdigit() else (1, part.lower()) for part in re.split(r"(\d+)", value) if part]
